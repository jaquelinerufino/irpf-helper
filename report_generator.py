import io
from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from document_checklist import DOCUMENT_CATEGORIES
from irpf_calc import format_currency

RECOMMENDED_LABELS = {
    "simplified": "Desconto Simplificado",
    "complete": "Deduções Completas",
    "equal": "Empate entre os modelos",
}


def _checklist_rows(checklist_state: List[Dict[str, Any]]) -> List[List[str]]:
    state_by_id = {item.get("id"): item for item in checklist_state}
    rows = []
    for category in DOCUMENT_CATEGORIES:
        item = state_by_id.get(category["id"], {})
        checked = "Sim" if item.get("checked") else "Não"
        file_names = item.get("fileNames") or []
        file_names_text = ", ".join(file_names) if file_names else "-"
        rows.append([category["group"], category["label"], checked, file_names_text])
    return rows


def generate_pdf_report(calc_result: Dict[str, Any], checklist_state: List[Dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("IRPF Helper — Resumo da Declaração", styles["Title"]))
    story.append(Paragraph(f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Renda e Comparação de Modelos", styles["Heading2"]))
    simplified = calc_result["simplified"]
    complete = calc_result["complete"]

    comparison_data = [
        ["", "Desconto Simplificado", "Deduções Completas"],
        ["Base tributável", format_currency(simplified["taxable_base"]), format_currency(complete["taxable_base"])],
        ["Imposto estimado", format_currency(simplified["tax_amount"]), format_currency(complete["tax_amount"])],
        ["Alíquota efetiva", f"{simplified['effective_rate']:.2f}%", f"{complete['effective_rate']:.2f}%"],
    ]
    table = Table(comparison_data, colWidths=[4 * cm, 6 * cm, 6 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f5bd9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 0.4 * cm))

    recommended_label = RECOMMENDED_LABELS.get(calc_result["recommended"], "")
    if calc_result["recommended"] != "equal":
        recommendation_text = (
            f"<b>{recommended_label}</b> compensa mais e economiza "
            f"{format_currency(calc_result['difference'])} em relação ao outro modelo."
        )
    else:
        recommendation_text = "Os dois modelos resultam no mesmo imposto estimado."
    story.append(Paragraph(recommendation_text, styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Organizador de Documentos", styles["Heading2"]))
    checklist_data = [["Grupo", "Documento", "Marcado", "Arquivo anexado"]] + _checklist_rows(checklist_state)
    checklist_table = Table(checklist_data, colWidths=[3 * cm, 7 * cm, 2.5 * cm, 3.5 * cm])
    checklist_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1ba97f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]
        )
    )
    story.append(checklist_table)
    story.append(Spacer(1, 0.8 * cm))

    story.append(
        Paragraph(
            "Este documento é uma estimativa gerada para apoio de planejamento. "
            "Não substitui a orientação de um contador nem a declaração oficial do IRPF.",
            styles["Italic"],
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(calc_result: Dict[str, Any], checklist_state: List[Dict[str, Any]]) -> bytes:
    workbook = Workbook()
    bold = Font(bold=True)

    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    summary_sheet.append(["IRPF Helper — Resumo da Declaração"])
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append([f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    summary_sheet.append([])
    summary_sheet.append(["Renda anual", format_currency(calc_result["annual_income"])])
    summary_sheet.append([])
    recommended_label = RECOMMENDED_LABELS.get(calc_result["recommended"], "")
    summary_sheet.append(["Modelo recomendado", recommended_label])
    summary_sheet.append(["Economia estimada", format_currency(calc_result["difference"])])
    for row in summary_sheet.iter_rows(min_row=4, max_row=4):
        for cell in row:
            cell.font = bold

    comparison_sheet = workbook.create_sheet("Comparação")
    comparison_sheet.append(["", "Desconto Simplificado", "Deduções Completas"])
    simplified = calc_result["simplified"]
    complete = calc_result["complete"]
    comparison_sheet.append(
        ["Base tributável", simplified["taxable_base"], complete["taxable_base"]]
    )
    comparison_sheet.append(["Imposto estimado", simplified["tax_amount"], complete["tax_amount"]])
    comparison_sheet.append(
        ["Alíquota efetiva (%)", simplified["effective_rate"], complete["effective_rate"]]
    )
    for cell in comparison_sheet[1]:
        cell.font = bold

    checklist_sheet = workbook.create_sheet("Checklist")
    checklist_sheet.append(["Grupo", "Documento", "Marcado", "Arquivo anexado"])
    for cell in checklist_sheet[1]:
        cell.font = bold
    for row in _checklist_rows(checklist_state):
        checklist_sheet.append(row)

    for sheet in (summary_sheet, comparison_sheet, checklist_sheet):
        for column_cells in sheet.columns:
            length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
