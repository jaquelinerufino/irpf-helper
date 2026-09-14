import io
from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .document_checklist import DOCUMENT_CATEGORIES
from .irpf_calc import format_currency

RECOMMENDED_LABELS = {
    "simplified": "Desconto Simplificado",
    "complete": "Deduções Completas",
    "equal": "Empate entre os modelos",
}

_FORM_FIELD_LABELS = {
    "salary": "Salário",
    "extraIncome": "Rendimentos extras",
    "deductions": "Deduções",
    "irrf": "IRRF retido",
    "inss": "Contribuição INSS",
}


def _format_balance(balance: float) -> str:
    if balance > 0:
        return f"{format_currency(balance)} (a pagar)"
    if balance < 0:
        return f"{format_currency(abs(balance))} (a restituir)"
    return f"{format_currency(0)} (quitado)"


def _declaration_rows(calculation: Dict[str, Any]) -> List[List[str]]:
    """Cada campo que o usuário preencheu, mapeado pra ficha da Receita onde
    ele entra — o resumo que a pessoa realmente usa na hora de declarar."""
    calculation = calculation or {}
    salary_monthly = float(calculation.get("salary") or 0)
    return [
        [
            "Salário bruto",
            f"{format_currency(salary_monthly)}/mês ({format_currency(salary_monthly * 12)}/ano)",
            "Rendimentos Tributáveis Recebidos de Pessoas Jurídicas",
        ],
        [
            "Rendimentos extras anuais",
            format_currency(float(calculation.get("extraIncome") or 0)),
            "Rendimentos Sujeitos à Tributação Exclusiva/Definitiva ou Isentos (conforme a natureza)",
        ],
        [
            "Contribuição INSS",
            format_currency(float(calculation.get("inss") or 0)),
            "Dedução na Ficha de Rendimentos de Pessoas Jurídicas",
        ],
        [
            "IRRF retido na fonte",
            format_currency(float(calculation.get("irrf") or 0)),
            "Imposto Retido na Fonte, na Ficha de Rendimentos",
        ],
        [
            "Deduções gerais",
            format_currency(float(calculation.get("deductions") or 0)),
            "Ficha Pagamentos Efetuados",
        ],
        [
            "Pensão alimentícia / outros descontos",
            format_currency(float(calculation.get("pension") or 0)),
            "Ficha Pagamentos Efetuados",
        ],
        [
            "Número de dependentes",
            str(int(calculation.get("dependents") or 0)),
            "Ficha Dependentes",
        ],
    ]


def _applied_items_rows(items: List[Dict[str, Any]]) -> List[List[str]]:
    """Itens específicos que o usuário confirmou durante a extração dos
    informes (ex.: cada recibo de saúde, cada rendimento de banco), com a
    classificação que ele escolheu quando havia ambiguidade."""
    rows = []
    for item in items or []:
        field_label = _FORM_FIELD_LABELS.get(item.get("form_field"), item.get("form_field") or "-")
        rows.append(
            [
                item.get("label") or "-",
                format_currency(float(item.get("value") or 0)),
                field_label,
                item.get("category") or "-",
            ]
        )
    return rows


def _checklist_rows(checklist_state: List[Dict[str, Any]]) -> List[List[str]]:
    state_by_id = {item.get("id"): item for item in checklist_state}
    rows = []
    for category in DOCUMENT_CATEGORIES:
        item = state_by_id.get(category["id"], {})
        checked = "Sim" if item.get("checked") else "Não"
        file_names = item.get("fileNames") or []
        file_names_text = ", ".join(file_names) if file_names else "-"
        rows.append([category["group"], category["label"], checked, file_names_text])
    return rows


def generate_pdf_report(
    calc_result: Dict[str, Any],
    checklist_state: List[Dict[str, Any]],
    calculation: Dict[str, Any] = None,
    items: List[Dict[str, Any]] = None,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("IRPF Helper — Resumo da Declaração", styles["Title"]))
    story.append(Paragraph(f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Renda e Comparação de Modelos", styles["Heading2"]))
    simplified = calc_result["simplified"]
    complete = calc_result["complete"]

    comparison_data = [
        ["", "Desconto Simplificado", "Deduções Completas"],
        ["Base tributável", format_currency(simplified["taxable_base"]), format_currency(complete["taxable_base"])],
        ["Imposto estimado", format_currency(simplified["tax_amount"]), format_currency(complete["tax_amount"])],
        ["Alíquota efetiva", f"{simplified['effective_rate']:.2f}%", f"{complete['effective_rate']:.2f}%"],
        ["Saldo (após IRRF retido)", _format_balance(simplified["balance"]), _format_balance(complete["balance"])],
    ]
    table = Table(comparison_data, colWidths=[4 * cm, 6 * cm, 6 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f5bd9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 0.4 * cm))

    recommended_label = RECOMMENDED_LABELS.get(calc_result["recommended"], "")
    if calc_result["recommended"] != "equal":
        recommendation_text = (
            f"<b>{recommended_label}</b> compensa mais e economiza "
            f"{format_currency(calc_result['difference'])} em relação ao outro modelo."
        )
    else:
        recommendation_text = "Os dois modelos resultam no mesmo imposto estimado."
    story.append(Paragraph(recommendation_text, styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Resumo para a Declaração", styles["Heading2"]))
    declaration_data = [["Campo", "Valor", "Onde declarar"]] + _declaration_rows(calculation)
    declaration_table = Table(declaration_data, colWidths=[4.5 * cm, 5.5 * cm, 6 * cm])
    declaration_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f5bd9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]
        )
    )
    story.append(declaration_table)
    story.append(Spacer(1, 0.6 * cm))

    item_rows = _applied_items_rows(items)
    if item_rows:
        story.append(Paragraph("Itens Detalhados dos Informes", styles["Heading2"]))
        small_style = styles["Normal"].clone("small")
        small_style.fontSize = 8
        item_rows = [[Paragraph(row[0], small_style), *row[1:]] for row in item_rows]
        items_data = [["Item", "Valor", "Campo", "Classificação"]] + item_rows
        items_table = Table(items_data, colWidths=[7 * cm, 3 * cm, 3 * cm, 3 * cm])
        items_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1ba97f")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
                ]
            )
        )
        story.append(items_table)
        story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Organizador de Documentos", styles["Heading2"]))
    checklist_data = [["Grupo", "Documento", "Marcado", "Arquivo anexado"]] + _checklist_rows(checklist_state)
    checklist_table = Table(checklist_data, colWidths=[3 * cm, 7 * cm, 2.5 * cm, 3.5 * cm])
    checklist_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1ba97f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe7f1")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]
        )
    )
    story.append(checklist_table)
    story.append(Spacer(1, 0.8 * cm))

    story.append(
        Paragraph(
            "Este documento é uma estimativa gerada para apoio de planejamento. "
            "Não substitui a orientação de um contador nem a declaração oficial do IRPF.",
            styles["Italic"],
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(
    calc_result: Dict[str, Any],
    checklist_state: List[Dict[str, Any]],
    calculation: Dict[str, Any] = None,
    items: List[Dict[str, Any]] = None,
) -> bytes:
    workbook = Workbook()
    bold = Font(bold=True)

    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    summary_sheet.append(["IRPF Helper — Resumo da Declaração"])
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append([f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    summary_sheet.append([])
    summary_sheet.append(["Renda anual", format_currency(calc_result["annual_income"])])
    summary_sheet.append([])
    recommended_label = RECOMMENDED_LABELS.get(calc_result["recommended"], "")
    summary_sheet.append(["Modelo recomendado", recommended_label])
    summary_sheet.append(["Economia estimada", format_currency(calc_result["difference"])])
    for row in summary_sheet.iter_rows(min_row=4, max_row=4):
        for cell in row:
            cell.font = bold

    comparison_sheet = workbook.create_sheet("Comparação")
    comparison_sheet.append(["", "Desconto Simplificado", "Deduções Completas"])
    simplified = calc_result["simplified"]
    complete = calc_result["complete"]
    comparison_sheet.append(
        ["Base tributável", simplified["taxable_base"], complete["taxable_base"]]
    )
    comparison_sheet.append(["Imposto estimado", simplified["tax_amount"], complete["tax_amount"]])
    comparison_sheet.append(
        ["Alíquota efetiva (%)", simplified["effective_rate"], complete["effective_rate"]]
    )
    comparison_sheet.append(
        ["Saldo (após IRRF retido)", _format_balance(simplified["balance"]), _format_balance(complete["balance"])]
    )
    for cell in comparison_sheet[1]:
        cell.font = bold

    declaration_sheet = workbook.create_sheet("Declaração")
    declaration_sheet.append(["Campo", "Valor", "Onde declarar"])
    for cell in declaration_sheet[1]:
        cell.font = bold
    for row in _declaration_rows(calculation):
        declaration_sheet.append(row)

    sheets = [summary_sheet, comparison_sheet, declaration_sheet]

    item_rows = _applied_items_rows(items)
    if item_rows:
        items_sheet = workbook.create_sheet("Itens Detalhados")
        items_sheet.append(["Item", "Valor", "Campo", "Classificação"])
        for cell in items_sheet[1]:
            cell.font = bold
        for row in item_rows:
            items_sheet.append(row)
        sheets.append(items_sheet)

    checklist_sheet = workbook.create_sheet("Checklist")
    checklist_sheet.append(["Grupo", "Documento", "Marcado", "Arquivo anexado"])
    for cell in checklist_sheet[1]:
        cell.font = bold
    for row in _checklist_rows(checklist_state):
        checklist_sheet.append(row)
    sheets.append(checklist_sheet)

    for sheet in sheets:
        for column_cells in sheet.columns:
            length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
