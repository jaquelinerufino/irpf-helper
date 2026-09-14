import io

from openpyxl import load_workbook

from irpf_calc import compare_models
from report_generator import (
    _applied_items_rows,
    _declaration_rows,
    generate_excel_report,
    generate_pdf_report,
)

_CALCULATION = {
    "salary": 8500,
    "extraIncome": 15000,
    "deductions": 1200,
    "inss": 1800,
    "dependents": 1,
    "pension": 0,
    "irrf": 900,
}

_ITEMS = [
    {
        "label": "ASS.ODONTO.COPARTICI - TITULAR - ODONTOPREV S/A",
        "value": 52.74,
        "form_field": "deductions",
        "category": "Despesa de saúde",
    },
    {
        "label": "Rendimentos tributáveis (Itaú)",
        "value": 5.01,
        "form_field": "extraIncome",
        "category": None,
    },
]


def test_declaration_rows_map_fields_to_fichas():
    rows = _declaration_rows(_CALCULATION)
    by_field = {row[0]: row for row in rows}

    assert "R$ 8.500,00/mês" in by_field["Salário bruto"][1]
    assert "R$ 102.000,00/ano" in by_field["Salário bruto"][1]
    assert "Rendimentos Tributáveis" in by_field["Salário bruto"][2]

    assert by_field["Rendimentos extras anuais"][1] == "R$ 15.000,00"
    assert by_field["IRRF retido na fonte"][1] == "R$ 900,00"
    assert by_field["Número de dependentes"][1] == "1"


def test_declaration_rows_defaults_to_zero_for_missing_fields():
    rows = _declaration_rows({})
    by_field = {row[0]: row for row in rows}
    assert by_field["Rendimentos extras anuais"][1] == "R$ 0,00"
    assert by_field["Número de dependentes"][1] == "0"


def test_applied_items_rows_maps_form_field_to_friendly_label():
    rows = _applied_items_rows(_ITEMS)
    assert rows[0][0] == "ASS.ODONTO.COPARTICI - TITULAR - ODONTOPREV S/A"
    assert rows[0][1] == "R$ 52,74"
    assert rows[0][2] == "Deduções"
    assert rows[0][3] == "Despesa de saúde"
    # sem categoria (campo não ambíguo) -> "-"
    assert rows[1][2] == "Rendimentos extras"
    assert rows[1][3] == "-"


def test_applied_items_rows_empty_when_no_items():
    assert _applied_items_rows([]) == []
    assert _applied_items_rows(None) == []


def test_generate_pdf_report_with_declaration_and_items():
    result = compare_models(_CALCULATION)
    pdf_bytes = generate_pdf_report(result, [], _CALCULATION, _ITEMS)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 0


def test_generate_pdf_report_without_calculation_or_items_still_works():
    result = compare_models({"salary": 5000})
    pdf_bytes = generate_pdf_report(result, [])
    assert pdf_bytes.startswith(b"%PDF")


def test_generate_excel_report_includes_declaration_and_items_sheets():
    result = compare_models(_CALCULATION)
    xlsx_bytes = generate_excel_report(result, [], _CALCULATION, _ITEMS)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))

    assert "Declaração" in workbook.sheetnames
    assert "Itens Detalhados" in workbook.sheetnames

    declaration_rows = list(workbook["Declaração"].iter_rows(values_only=True))
    assert declaration_rows[0] == ("Campo", "Valor", "Onde declarar")

    item_rows = list(workbook["Itens Detalhados"].iter_rows(values_only=True))
    assert item_rows[1][0] == "ASS.ODONTO.COPARTICI - TITULAR - ODONTOPREV S/A"


def test_generate_excel_report_omits_items_sheet_when_no_items():
    result = compare_models(_CALCULATION)
    xlsx_bytes = generate_excel_report(result, [], _CALCULATION, [])
    workbook = load_workbook(io.BytesIO(xlsx_bytes))

    assert "Declaração" in workbook.sheetnames
    assert "Itens Detalhados" not in workbook.sheetnames
